from decimal import Decimal
import json
from pathlib import Path
import warnings

import pytest

from exporters import export_results
from models import (
    AggregationStatistic,
    DocumentMetadata,
    DatasetCharacteristics,
    EvidenceRecord,
    EvidenceType,
    ExtractionProvenance,
    FAIRagroApplicationDataFitnessModel,
    IdentifierObjectType,
    IdentifierStatus,
    MetricContext,
    MetricScope,
    MetricValueKind,
    QuantitativeMetricValue,
    RelatedResource,
    RelatedResourceType,
    RunPurpose,
    ValidationAndDiagnostics,
    ValidationMetricRecord,
)
from quantitative_values import enrich_quantitative_values, parse_quantitative_value
from semantic_checks import validate_record
from source_bundle import SourceBundle, _build_qualitative_guardrail_hints, _portable_relpath


def _bundle(**overrides) -> SourceBundle:
    values = {
        "package_dir": "/tmp/package",
        "source_filename": "paper.pdf",
        "source_sha256": "a" * 64,
        "source_fidelity_status": "pass",
        "source_fidelity_initial_status": "pass",
        "source_fidelity_summary": {},
        "unresolved_items": [],
        "clean_markdown_path": "/tmp/paper.md",
        "text_for_llm": "source",
        "bundle_sha256": "b" * 64,
        "item_counts": {},
    }
    values.update(overrides)
    return SourceBundle(**values)


def _provenance() -> ExtractionProvenance:
    return ExtractionProvenance(
        schema_version="fairagro-dffp-v3.5.1",
        canonical_source_sha256="a" * 64,
        source_bundle_sha256="b" * 64,
        run_purpose=RunPurpose.test,
    )


@pytest.mark.parametrize(
    "display",
    [
        "2.8-5.1 days (median 3.9 days)",
        "median 3.9 days (range 2.8-5.1 days)",
    ],
)
def test_central_statistic_and_range_are_preserved_together(display: str):
    value = parse_quantitative_value(display, metric_name="MAE", unit="days")
    assert value.kind.value == "estimate_with_range"
    assert value.numeric_value == Decimal("3.9")
    assert value.lower_bound == Decimal("2.8")
    assert value.upper_bound == Decimal("5.1")
    assert value.aggregation == AggregationStatistic.median
    assert value.as_reported == display


def test_structured_decimals_serialize_as_json_numbers():
    value = parse_quantitative_value("0.900 (0.888-0.920)", metric_name="PICP")
    payload = value.model_dump(mode="json")
    assert isinstance(payload["numeric_value"], float)
    assert payload["numeric_value"] == 0.9
    assert payload["lower_bound"] == 0.888


def test_incomplete_estimate_with_range_is_safely_downgraded_before_sdk_parse():
    value = QuantitativeMetricValue.model_validate({
        "kind": "estimate_with_range",
        "as_reported": "2.8-5.1",
        "numeric_value": None,
        "lower_bound": 2.8,
        "upper_bound": 5.1,
        "nominal_level": None,
        "unit_code": "d",
        "aggregation": "unspecified",
        "range_basis": "annual range",
        "points": None,
        "uncertainty_components": None,
    })
    assert value.kind == MetricValueKind.range
    assert value.numeric_value is None
    assert value.lower_bound == Decimal("2.8")
    assert value.upper_bound == Decimal("5.1")


def test_incomplete_estimate_without_bounds_becomes_scalar():
    value = QuantitativeMetricValue.model_validate({
        "kind": "estimate_with_range",
        "as_reported": "3.9",
        "numeric_value": 3.9,
        "lower_bound": None,
        "upper_bound": None,
    })
    assert value.kind == MetricValueKind.scalar
    assert value.numeric_value == Decimal("3.9")


def test_direct_evidence_recovers_central_value_omitted_from_range_display():
    metric = ValidationMetricRecord(
        name="MAE",
        value_or_summary="2.8-5.1",
        unit="days",
        quantitative_value=QuantitativeMetricValue(
            kind=MetricValueKind.range,
            as_reported="2.8-5.1",
            lower_bound=Decimal("2.8"),
            upper_bound=Decimal("5.1"),
            range_basis="annual range",
        ),
        context=MetricContext.out_of_sample_validation,
        scope=MetricScope.crop_phase,
        scope_label="Heading phase",
        evidence=[EvidenceRecord(
            claim="Heading MAE is reported.",
            evidence_type=EvidenceType.explicit,
            source_text="The heading phase has an MAE of 2.8-5.1 days (median 3.9 days).",
        )],
    )
    record = FAIRagroApplicationDataFitnessModel(
        validation_and_diagnostics=ValidationAndDiagnostics(validation_metrics=[metric])
    )
    enrich_quantitative_values(record)
    assert metric.value_or_summary == "2.8-5.1 (median 3.9)"
    assert metric.quantitative_value.kind == MetricValueKind.estimate_with_range
    assert metric.quantitative_value.numeric_value == Decimal("3.9")
    assert metric.quantitative_value.range_basis == "annual range"


def test_windows_manifest_path_is_portable():
    assert _portable_relpath(r"docling\document.clean.md") == Path("docling/document.clean.md")


def test_numbered_limitation_section_becomes_guardrail_checklist():
    markdown = """# Paper
## 4 Limitations
1. Validation is conditional on prior filtering.
2. Temperature uncertainty is not propagated.
## 5 Availability
Data are available.
"""
    hints = _build_qualitative_guardrail_hints(markdown)
    assert len(hints) == 1
    assert hints[0]["source_section"] == "4 Limitations"
    assert [x["number"] for x in hints[0]["numbered_items"]] == [1, 2]


def test_source_explicit_identifier_without_evidence_is_an_error():
    record = FAIRagroApplicationDataFitnessModel(
        document_metadata=DocumentMetadata(related_resources=[RelatedResource(
            name="Example dataset",
            identifier="https://doi.org/10.1234/example",
            resource_type=RelatedResourceType.dataset,
            identifier_object_type=IdentifierObjectType.dataset,
            identifier_status=IdentifierStatus.source_explicit,
        )]),
        extraction_provenance=_provenance(),
    )
    issues = validate_record(record, _bundle())
    issue = next(x for x in issues if x.code == "EXPLICIT_RESOURCE_IDENTIFIER_LACKS_EVIDENCE")
    assert issue.severity == "error"


def test_source_hash_mismatch_is_an_error():
    record = FAIRagroApplicationDataFitnessModel(extraction_provenance=_provenance())
    record.extraction_provenance.canonical_source_sha256 = "c" * 64
    issues = validate_record(record, _bundle())
    assert any(x.code == "CANONICAL_SOURCE_HASH_MISMATCH" and x.severity == "error" for x in issues)


def test_validated_export_includes_schema_and_ro_crate_descriptor(tmp_path: Path):
    evidence = EvidenceRecord(
        claim="Dataset identifier is stated.",
        evidence_type=EvidenceType.explicit,
        source_text="Data: https://doi.org/10.1234/example",
    )
    record = FAIRagroApplicationDataFitnessModel(
        document_metadata=DocumentMetadata(
            title="Example data paper",
            related_resources=[RelatedResource(
                name="Example dataset",
                identifier="https://doi.org/10.1234/example",
                resource_type=RelatedResourceType.dataset,
                identifier_object_type=IdentifierObjectType.dataset,
                identifier_status=IdentifierStatus.source_explicit,
                evidence=[evidence],
            )],
        ),
        extraction_provenance=_provenance(),
    )
    paths = export_results([record], {"semantic_checks": []}, output_dir=str(tmp_path))
    assert Path(paths.schema_path).exists()
    assert paths.ro_crate_metadata_path is not None
    crate = json.loads(Path(paths.ro_crate_metadata_path).read_text(encoding="utf-8"))
    assert crate["@context"] == "https://w3id.org/ro/crate/1.1/context"
    assert any(x.get("@id") == "application_matrix.json" for x in crate["@graph"])

    jsonschema = pytest.importorskip("jsonschema")
    payload = json.loads(Path(paths.json_path).read_text(encoding="utf-8"))
    schema = json.loads(Path(paths.schema_path).read_text(encoding="utf-8"))
    jsonschema.validate(payload, schema)


def test_duplicate_prose_window_with_noisy_target_does_not_create_false_missing_metric():
    evidence = EvidenceRecord(
        claim="Winter-wheat phase summaries are reported.",
        evidence_type=EvidenceType.explicit,
        source_section="3.3.2 Accuracy across phenological phases",
        source_text="Across seven winter wheat phases, median MAE is 5.7 and median RMSE is 7.1 days.",
    )
    metrics = [
        ValidationMetricRecord(
            name="MAE",
            value_or_summary="2.6-9.5 days (median 5.7 days)",
            unit="days",
            quantitative_value=parse_quantitative_value("2.6-9.5 days (median 5.7 days)", metric_name="MAE", unit="days"),
            context=MetricContext.out_of_sample_validation,
            scope=MetricScope.crop_phase,
            scope_label="seven winter wheat phases",
            evidence=[evidence],
        ),
        ValidationMetricRecord(
            name="RMSE",
            value_or_summary="3.4-11.6 days (median 7.1 days)",
            unit="days",
            quantitative_value=parse_quantitative_value("3.4-11.6 days (median 7.1 days)", metric_name="RMSE", unit="days"),
            context=MetricContext.out_of_sample_validation,
            scope=MetricScope.crop_phase,
            scope_label="seven winter wheat phases",
            evidence=[evidence],
        ),
    ]
    common = {
        "kind": "prose_metric_candidate",
        "metric_names": ["MAE", "RMSE"],
        "metric_keys": ["mae", "rmse"],
        "result_metric_keys": ["mae", "rmse"],
        "quantitative_pairs": [
            {"metric_name": "MAE", "metric_key": "mae", "value": "5.7"},
            {"metric_name": "RMSE", "metric_key": "rmse", "value": "7.1"},
        ],
        "summary_signal": True,
        "quality_context": True,
        "source_section": "3.3.2 Accuracy across phenological phases",
    }
    hints = [
        {**common, "target_hint": "every phase and year"},
        {**common, "target_hint": "seven winter wheat phases"},
    ]
    record = FAIRagroApplicationDataFitnessModel(
        validation_and_diagnostics=ValidationAndDiagnostics(validation_metrics=metrics),
        extraction_provenance=_provenance(),
    )
    issues = validate_record(record, _bundle(metric_candidate_hints=hints))
    assert not any(x.code == "SUMMARY_PROSE_METRIC_UNREPRESENTED" for x in issues)


def test_excel_export_chunks_large_nested_fields_without_cell_limit_warning(tmp_path: Path):
    record = FAIRagroApplicationDataFitnessModel(
        dataset_characteristics=DatasetCharacteristics(description="x" * 70000),
        extraction_provenance=_provenance(),
    )
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        paths = export_results([record], {"semantic_checks": []}, output_dir=str(tmp_path))
    assert not any("Cell contents too long" in str(item.message) for item in caught)

    import openpyxl
    workbook = openpyxl.load_workbook(paths.excel_path, read_only=True)
    assert "record_fields" in workbook.sheetnames
    assert "evidence" in workbook.sheetnames
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            assert all(not isinstance(cell.value, str) or len(cell.value) <= 32767 for cell in row)
