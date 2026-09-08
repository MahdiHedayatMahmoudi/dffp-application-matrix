"""Export DFFP records and technical provenance independently of Streamlit."""

from __future__ import annotations

import json
import hashlib
import tempfile
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Any

import pandas as pd

from models import FAIRagroApplicationDataFitnessModel


@dataclass(frozen=True)
class ExportPaths:
    output_dir: str
    json_path: str
    manifest_path: str
    validation_report_path: str
    excel_path: str
    html_path: str
    interactive_html_path: str
    schema_path: str
    ro_crate_metadata_path: str | None


def record_to_dict(record: FAIRagroApplicationDataFitnessModel) -> dict[str, Any]:
    return record.model_dump(mode="json", by_alias=True, exclude_none=True)

def _portable_relative_path(
    value: str | None,
    *,
    package_dir: str | None = None,
) -> str | None:
    """Convert a runtime filesystem path into a portable output locator.

    Absolute workstation paths are needed internally while the pipeline runs,
    but they should not be serialized into canonical output artifacts.
    """
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return text

    # Handle Windows paths correctly even if this code is inspected/tested
    # on another operating system.
    is_windows = (
        len(text) >= 3
        and text[1] == ":"
        and text[2] in "\\/"
    ) or "\\" in text

    path_cls = PureWindowsPath if is_windows else PurePosixPath
    path = path_cls(text)

    # Canonical exports may be revalidated repeatedly.  Once a path is already
    # relative/portable, preserve it exactly (apart from separator normalization)
    # instead of collapsing another directory level on every export pass.
    if not path.is_absolute():
        return path.as_posix()

    # If the file lives inside the scientific source package,
    # keep its path relative to that package.
    if package_dir:
        package_text = str(package_dir).strip()

        package_is_windows = (
            len(package_text) >= 3
            and package_text[1] == ":"
            and package_text[2] in "\\/"
        ) or "\\" in package_text

        if package_is_windows == is_windows:
            package = path_cls(package_text)

            try:
                return path.relative_to(package).as_posix()
            except ValueError:
                pass

    # If it is outside the package, retain only the final name.
    return path.name or text


def _portable_manifest(manifest: dict[str, Any]) -> dict[str, Any]:
    """Return a portable copy of an extraction manifest.

    The original runtime manifest is left unchanged.
    """
    payload = deepcopy(manifest)

    source_package = payload.get("source_package")

    if isinstance(source_package, dict):
        # Important: preserve the original absolute package path temporarily,
        # because we need it to calculate relative child paths.
        original_package_dir = (
            (manifest.get("source_package") or {}).get("package_dir")
        )

        if original_package_dir:
            source_package["package_dir"] = _portable_relative_path(
                original_package_dir
            )

        clean_markdown_path = source_package.get("clean_markdown_path")

        if clean_markdown_path:
            source_package["clean_markdown_path"] = _portable_relative_path(
                clean_markdown_path,
                package_dir=original_package_dir,
            )

    # source_bundle_path is currently another absolute path generated
    # by pipeline.py.
    if payload.get("source_bundle_path"):
        original_package_dir = (
            (manifest.get("source_package") or {}).get("package_dir")
        )

        payload["source_bundle_path"] = _portable_relative_path(
            payload["source_bundle_path"],
            package_dir=original_package_dir,
        )

    # Optional but recommended:
    # response IDs are provider-side execution identifiers. They are not
    # credentials, but the scientific output does not need them.
    for key in (
        "extraction_run",
        "final_structured_run",
        "revalidation_run",
    ):
        run = payload.get(key)

        if isinstance(run, dict):
            run.pop("response_id", None)

    repair_runs = payload.get("semantic_repair_runs")

    if isinstance(repair_runs, list):
        for run in repair_runs:
            if isinstance(run, dict):
                run.pop("response_id", None)

    # Make the policy explicit in the output itself.
    payload["output_portability"] = {
        "absolute_local_paths_included": False,
        "path_policy": "source-package-relative-or-basename",
    }

    return payload



def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _application_matrix_schema() -> dict[str, Any]:
    """Return the public JSON Schema for the actual top-level array artifact."""

    record_schema = FAIRagroApplicationDataFitnessModel.model_json_schema(by_alias=True)
    definitions = record_schema.pop("$defs", {})
    schema = {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "title": "FAIRagro DFFP application matrix",
        "description": "One or more source-grounded FAIRagro DFFP evidence records.",
        "type": "array",
        "items": record_schema,
    }
    if definitions:
        schema["$defs"] = definitions
    return schema


def _write_ro_crate_metadata(
    path: Path,
    *,
    json_path: Path,
    schema_path: Path,
    records: list[FAIRagroApplicationDataFitnessModel],
) -> None:
    """Write a minimal valid RO-Crate JSON-LD descriptor for the output directory.

    The application matrix remains ordinary validated JSON. RO-Crate supplies the
    packaging graph and links it to the public schema and canonical-source identity.
    """

    first = records[0] if records else None
    metadata = first.document_metadata if first else None
    provenance = first.extraction_provenance if first else None
    source_hash = provenance.canonical_source_sha256 if provenance else None
    source_id = f"urn:sha256:{source_hash}" if source_hash else "#canonical-source"
    root_name = metadata.title if metadata and metadata.title else "FAIRagro DFFP evidence package"
    graph: list[dict[str, Any]] = [
        {
            "@id": "ro-crate-metadata.json",
            "@type": "CreativeWork",
            "about": {"@id": "./"},
            "conformsTo": {"@id": "https://w3id.org/ro/crate/1.1"},
        },
        {
            "@id": "./",
            "@type": "Dataset",
            "name": root_name,
            "hasPart": [
                {"@id": json_path.name},
                {"@id": schema_path.name},
                {"@id": "extraction_manifest.json"},
                {"@id": "validation_report.json"},
            ],
            "source": {"@id": source_id},
        },
        {
            "@id": json_path.name,
            "@type": "File",
            "name": "FAIRagro DFFP application matrix",
            "encodingFormat": "application/json",
            "sha256": _sha256_file(json_path),
            "conformsTo": {"@id": schema_path.name},
            "about": {"@id": source_id},
        },
        {
            "@id": schema_path.name,
            "@type": "File",
            "name": "Application matrix JSON Schema",
            "encodingFormat": "application/schema+json",
        },
        {"@id": "extraction_manifest.json", "@type": "File", "encodingFormat": "application/json"},
        {"@id": "validation_report.json", "@type": "File", "encodingFormat": "application/json"},
        {
            "@id": source_id,
            "@type": "ScholarlyArticle",
            "name": metadata.title if metadata and metadata.title else "Canonical source document",
            "identifier": metadata.doi if metadata and metadata.doi else source_hash,
            "sha256": source_hash,
        },
    ]
    path.write_text(
        json.dumps({"@context": "https://w3id.org/ro/crate/1.1/context", "@graph": graph},
                   ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _format(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2, default=str)
    return str(value)


def _record_summary_rows(records: list[FAIRagroApplicationDataFitnessModel]) -> list[dict[str, Any]]:
    """Create a review-friendly record index instead of one giant nested-JSON row."""

    rows: list[dict[str, Any]] = []
    for record_index, record in enumerate(records):
        metadata = record.document_metadata
        dataset = record.dataset_characteristics
        spatial = dataset.spatial if dataset else None
        temporal = dataset.temporal if dataset else None
        diagnostics = record.validation_and_diagnostics
        provenance = record.extraction_provenance
        limitations = record.limitations_and_risks
        limitation_count = 0
        if limitations:
            limitation_count = sum(len(value or []) for value in (
                limitations.known_limitations,
                limitations.risk_of_misuse,
                limitations.bias_sources,
                limitations.extrapolation_limits,
            ))
        rows.append({
            "record_index": record_index,
            "document_title": metadata.title if metadata else None,
            "document_doi": metadata.doi if metadata else None,
            "document_status": metadata.document_status.value if metadata else None,
            "source_filename": metadata.source if metadata else None,
            "canonical_source_sha256": provenance.canonical_source_sha256 if provenance else None,
            "schema_version": provenance.schema_version if provenance else None,
            "run_purpose": provenance.run_purpose.value if provenance else None,
            "review_status": provenance.review_status.value if provenance else None,
            "dataset_name": dataset.dataset_name if dataset else None,
            "geographic_scope": dataset.geographic_scope if dataset else None,
            "native_resolution": spatial.native_resolution if spatial else None,
            "temporal_extent": temporal.temporal_extent if temporal else None,
            "validation_metric_count": len(diagnostics.validation_metrics or []) if diagnostics else 0,
            "tuning_parameter_count": len(diagnostics.tuning_parameters or []) if diagnostics else 0,
            "related_resource_count": len(metadata.related_resources or []) if metadata else 0,
            "limitation_and_risk_count": limitation_count,
        })
    return rows


def _record_field_rows(records: list[FAIRagroApplicationDataFitnessModel]) -> list[dict[str, Any]]:
    """Flatten every scalar into auditable path/value rows without Excel-sized blobs."""

    rows: list[dict[str, Any]] = []
    chunk_size = 4000

    def walk(value: Any, record_index: int, path: str) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                walk(child, record_index, f"{path}.{key}" if path else key)
            return
        if isinstance(value, list):
            for item_index, child in enumerate(value):
                walk(child, record_index, f"{path}[{item_index}]")
            return
        rendered = "" if value is None else str(value)
        chunks = [rendered[i:i + chunk_size] for i in range(0, len(rendered), chunk_size)] or [""]
        for chunk_index, chunk in enumerate(chunks, start=1):
            rows.append({
                "record_index": record_index,
                "field_path": path,
                "value_type": type(value).__name__,
                "chunk_index": chunk_index,
                "chunk_count": len(chunks),
                "value": chunk,
            })

    for record_index, record in enumerate(records):
        walk(record_to_dict(record), record_index, "")
    return rows


def _evidence_rows(records: list[FAIRagroApplicationDataFitnessModel]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    def walk(value: Any, record_index: int, path: str) -> None:
        if isinstance(value, dict):
            if "claim" in value and "evidence_type" in value and "review_status" in value:
                rows.append({
                    "record_index": record_index,
                    "field_path": path,
                    "claim": value.get("claim"),
                    "evidence_type": value.get("evidence_type"),
                    "source_section": value.get("source_section"),
                    "source_location": value.get("source_location"),
                    "source_page": value.get("source_page"),
                    "source_item_id": value.get("source_item_id"),
                    "source_modality": value.get("source_modality"),
                    "representation_method": value.get("representation_method"),
                    "source_text": value.get("source_text"),
                    "review_status": value.get("review_status"),
                })
            for key, child in value.items():
                walk(child, record_index, f"{path}.{key}" if path else key)
        elif isinstance(value, list):
            for item_index, child in enumerate(value):
                walk(child, record_index, f"{path}[{item_index}]")

    for record_index, record in enumerate(records):
        walk(record_to_dict(record), record_index, "")
    return rows


def _excel_safe_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Protect Excel's 32,767-character cell limit; canonical JSON stays lossless."""

    def safe(value: Any) -> Any:
        if isinstance(value, str) and len(value) > 30000:
            return value[:29950] + " … [truncated in Excel; see application_matrix JSON]"
        return value

    return df.map(safe)


def _style_excel_sheets(writer: pd.ExcelWriter) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="top")
        for column_cells in worksheet.iter_cols(min_row=1, max_row=min(worksheet.max_row, 200)):
            maximum = max((len(str(cell.value or "")) for cell in column_cells), default=0)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(maximum + 2, 10), 48)


def _metric_rows(records: list[FAIRagroApplicationDataFitnessModel]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record_index, record in enumerate(records):
        metadata = record.document_metadata
        diagnostics = record.validation_and_diagnostics
        for row_type, values in (
            ("validation_metric", list(diagnostics.validation_metrics or []) if diagnostics else []),
            ("tuning_parameter", list(diagnostics.tuning_parameters or []) if diagnostics else []),
        ):
            for metric in values:
                q = metric.quantitative_value
                evidence = list(metric.evidence or [])
                rows.append({
                    "record_index": record_index,
                    "document_title": metadata.title if metadata else None,
                    "document_doi": metadata.doi if metadata else None,
                    "row_type": row_type,
                    "name": metric.name,
                    "value_or_summary": metric.value_or_summary,
                    "unit_as_reported": metric.unit,
                    "value_kind": _format(q.kind.value if q else None),
                    "numeric_value": _format(q.numeric_value if q else None),
                    "lower_bound": _format(q.lower_bound if q else None),
                    "upper_bound": _format(q.upper_bound if q else None),
                    "nominal_level": _format(q.nominal_level if q else None),
                    "unit_code": q.unit_code if q else None,
                    "aggregation": q.aggregation.value if q else None,
                    "range_basis": q.range_basis if q else None,
                    "points": _format(q.model_dump(mode="json").get("points") if q and q.points else None),
                    "uncertainty_components": _format(
                        q.model_dump(mode="json").get("uncertainty_components")
                        if q and q.uncertainty_components else None
                    ),
                    "context": _format(getattr(metric, "context", None).value if getattr(metric, "context", None) else None),
                    "scope": _format(metric.scope.value),
                    "scope_label": metric.scope_label,
                    "evaluation_support": _format(
                        getattr(metric, "evaluation_support", None).value
                        if getattr(metric, "evaluation_support", None) else None
                    ),
                    "evaluation_population": getattr(metric, "evaluation_population", None),
                    "evidence_count": len(evidence),
                    "source_item_ids": "; ".join(sorted({x.source_item_id for x in evidence if x.source_item_id})),
                    "source_locations": "; ".join(sorted({x.source_location for x in evidence if x.source_location})),
                })
    return rows


def _resource_rows(records: list[FAIRagroApplicationDataFitnessModel]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record_index, record in enumerate(records):
        metadata = record.document_metadata
        for resource in list(metadata.related_resources or []) if metadata else []:
            rows.append({
                "record_index": record_index,
                "document_title": metadata.title,
                "name": resource.name,
                "identifier": resource.identifier,
                "resource_type": resource.resource_type.value,
                "identifier_object_type": resource.identifier_object_type.value,
                "identifier_status": resource.identifier_status.value,
                "identifier_note": resource.identifier_note,
                "relation": resource.relation,
                "version": resource.version,
                "is_concept_identifier": resource.is_concept_identifier,
                "evidence_count": len(resource.evidence or []),
            })
    return rows


def _quality_summary(records: list[FAIRagroApplicationDataFitnessModel], issues: list[dict[str, Any]]) -> dict[str, Any]:
    metrics = _metric_rows(records)
    validation = [x for x in metrics if x["row_type"] == "validation_metric"]
    resources = _resource_rows(records)
    return {
        "record_count": len(records),
        "canonical_validation_metric_count": len(validation),
        "structured_numeric_metric_count": sum(1 for x in validation if x["value_kind"] not in {"", "text_summary"}),
        "text_summary_metric_count": sum(1 for x in validation if x["value_kind"] == "text_summary"),
        "metric_target_count": len({(x["scope"], x["scope_label"]) for x in validation}),
        "metrics_with_evidence_count": sum(1 for x in validation if x["evidence_count"] > 0),
        "related_resource_count": len(resources),
        "verified_resource_identifier_count": sum(
            1 for x in resources if x["identifier_status"] == "source_explicit" and x["identifier_object_type"] != "unknown"
        ),
        "semantic_issue_counts": {
            severity: sum(1 for x in issues if x.get("severity") == severity)
            for severity in ("error", "warning", "info")
        },
    }


def _write_static_html(tables: list[tuple[str, pd.DataFrame]], path: Path, title: str) -> None:
    sections = []
    for heading, df in tables:
        sections.append(f"<h2>{heading}</h2>" + df.to_html(index=False, escape=True))
    table = "".join(sections)
    path.write_text(
        f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{title}</title>
<style>body{{font-family:Arial,sans-serif;margin:24px}}.wrap{{overflow-x:auto}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ddd;padding:8px;vertical-align:top;white-space:pre-wrap}}th{{font-weight:700}}</style>
</head><body><h1>{title}</h1><div class="wrap">{table}</div></body></html>""",
        encoding="utf-8",
    )


def _write_interactive_html(payload: dict[str, Any], path: Path) -> None:
    data = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    html = """<!doctype html><html lang="en"><head><meta charset="utf-8"><title>FAIRagro DFFP Explorer</title>
<style>body{font-family:Arial,sans-serif;margin:24px;max-width:1300px}details{border:1px solid #ddd;border-radius:8px;margin:8px 0;padding:10px}summary{cursor:pointer;font-weight:700}pre{white-space:pre-wrap;overflow-wrap:anywhere}</style></head>
<body><h1>FAIRagro DFFP Explorer</h1><div id="root"></div><script id="data" type="application/json">__DATA__</script>
<script>const data=JSON.parse(document.getElementById('data').textContent);const root=document.getElementById('root');function add(t,v){const d=document.createElement('details');const s=document.createElement('summary');s.textContent=t;const p=document.createElement('pre');p.textContent=JSON.stringify(v,null,2);d.appendChild(s);d.appendChild(p);root.appendChild(d);}Object.entries(data).forEach(([k,v])=>add(k,v));</script></body></html>""".replace("__DATA__", data)
    path.write_text(html, encoding="utf-8")


def export_results(
    records: list[FAIRagroApplicationDataFitnessModel],
    manifest: dict[str, Any],
    *,
    output_dir: str | None = None,
    json_filename: str = "application_matrix.json",
) -> ExportPaths:
    out = Path(output_dir or tempfile.mkdtemp(prefix="fairagro_dffp_"))
    out.mkdir(parents=True, exist_ok=True)

    json_path = out / json_filename
    manifest_path = out / "extraction_manifest.json"
    validation_path = out / "validation_report.json"
    excel_path = out / "application_matrix.xlsx"
    html_path = out / "application_matrix.html"
    interactive_path = out / "application_matrix_interactive.html"
    schema_path = out / "application_matrix.schema.json"
    ro_crate_path = out / "ro-crate-metadata.json"

    record_dicts = [record_to_dict(r) for r in records]
    json_path.write_text(json.dumps(record_dicts, ensure_ascii=False, indent=2), encoding="utf-8")
    schema_path.write_text(
        json.dumps(_application_matrix_schema(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    issues = list(manifest.get("semantic_checks", []))
    manifest_payload = _portable_manifest(manifest)
    manifest_payload["quality_summary"] = _quality_summary(records, issues)
    manifest_path.write_text(json.dumps(manifest_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    validation_path.write_text(
        json.dumps(issues, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    df = pd.DataFrame(_record_summary_rows(records))
    field_df = pd.DataFrame(_record_field_rows(records))
    metric_df = pd.DataFrame(_metric_rows(records))
    resource_df = pd.DataFrame(_resource_rows(records))
    evidence_df = pd.DataFrame(_evidence_rows(records))
    checks_df = pd.DataFrame(issues)
    excel_tables = [df, field_df, metric_df, resource_df, evidence_df, checks_df]
    df, field_df, metric_df, resource_df, evidence_df, checks_df = [
        _excel_safe_dataframe(table) for table in excel_tables
    ]
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="records", index=False)
        field_df.to_excel(writer, sheet_name="record_fields", index=False)
        metric_df.to_excel(writer, sheet_name="quantitative_metrics", index=False)
        resource_df.to_excel(writer, sheet_name="related_resources", index=False)
        evidence_df.to_excel(writer, sheet_name="evidence", index=False)
        checks_df.to_excel(writer, sheet_name="semantic_checks", index=False)
        _style_excel_sheets(writer)
    _write_static_html(
        [("Records", df), ("Quantitative metrics", metric_df), ("Related resources", resource_df),
         ("Evidence", evidence_df), ("Semantic checks", checks_df)],
        html_path,
        "FAIRagro Application Matrix",
    )
    _write_interactive_html({"records": record_dicts, "manifest": manifest_payload}, interactive_path)

    ro_crate_metadata_path: str | None = None
    if json_path.name == "application_matrix.json":
        _write_ro_crate_metadata(
            ro_crate_path,
            json_path=json_path,
            schema_path=schema_path,
            records=records,
        )
        ro_crate_metadata_path = str(ro_crate_path)

    return ExportPaths(
        output_dir=str(out),
        json_path=str(json_path),
        manifest_path=str(manifest_path),
        validation_report_path=str(validation_path),
        excel_path=str(excel_path),
        html_path=str(html_path),
        interactive_html_path=str(interactive_path),
        schema_path=str(schema_path),
        ro_crate_metadata_path=ro_crate_metadata_path,
    )
